import pandas as pd
import openpyxl
import sys
import numpy as np

def manual_stratified_split(df, train_frac=0.7, val_frac=0.15, random_state=42):
    np.random.seed(random_state)
    train_list, val_list, test_list = [], [], []
    
    for label, group in df.groupby('labels'):
        group = group.sample(frac=1, random_state=random_state).reset_index(drop=True)
        n = len(group)
        train_end = int(n * train_frac)
        val_end = train_end + int(n * val_frac)
        
        train_list.append(group.iloc[:train_end])
        val_list.append(group.iloc[train_end:val_end])
        test_list.append(group.iloc[val_end:])
        
    train_df = pd.concat(train_list).sample(frac=1, random_state=random_state).reset_index(drop=True)
    val_df = pd.concat(val_list).sample(frac=1, random_state=random_state).reset_index(drop=True)
    test_df = pd.concat(test_list).sample(frac=1, random_state=random_state).reset_index(drop=True)
    
    return train_df, val_df, test_df

def main():
    print("Loading dataset.xlsx...")
    try:
        wb = openpyxl.load_workbook("dataset.xlsx")
        ws = wb.active
        rows = []
        for r in range(2, ws.max_row + 1):
            text  = ws.cell(row=r, column=1).value
            label = ws.cell(row=r, column=2).value
            if text and label:
                rows.append({
                    "posts":  str(text).strip(),
                    "labels": int(label),
                })
        df = pd.DataFrame(rows)
    except Exception as e:
        print(f"Error loading dataset: {e}")
        sys.exit(1)
        
    print(f"Total rows loaded: {len(df)}")
    
    print("Performing stratified split (70% Train, 15% Val, 15% Test)...")
    train_df, val_df, test_df = manual_stratified_split(df)
    
    print(f"\nTrain size: {len(train_df)} ({len(train_df)/len(df)*100:.1f}%)")
    print(f"Val size  : {len(val_df)} ({len(val_df)/len(df)*100:.1f}%)")
    print(f"Test size : {len(test_df)} ({len(test_df)/len(df)*100:.1f}%)")
    
    print("\nLabel distribution in Train:")
    print(train_df['labels'].value_counts().sort_index())
    
    print("\nLabel distribution in Val:")
    print(val_df['labels'].value_counts().sort_index())
    
    print("\nLabel distribution in Test:")
    print(test_df['labels'].value_counts().sort_index())

    print("\nSaving splits to Excel files...")
    train_df.to_excel("train.xlsx", index=False)
    val_df.to_excel("val.xlsx", index=False)
    test_df.to_excel("test.xlsx", index=False)
    
    print("Done! Files saved: train.xlsx, val.xlsx, test.xlsx")

if __name__ == "__main__":
    main()
